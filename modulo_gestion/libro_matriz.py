import io
import datetime
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from modulo_base_datos.models import Carrera, Alumno, Cursada, PlanEstudio, Materia, Evaluacion

def generar_libro_matriz_excel(carrera_obj):
    """
    Genera el archivo Excel oficial del Libro Matriz para una carrera específica
    según la estructura reglamentaria del ISFT N° 188.
    Incluye:
    - Resumen general con estudiantes, cohorte, estado de egreso y legajo.
    - Planillas de calificaciones por cada año de cursada.
    - Planilla consolidada de egresados.
    Retorna un objeto BytesIO con el libro de cálculo generado.
    """
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Estilos de formato institucional
    font_title = Font(name='Arial', size=13, bold=True, color='1E3A8A')
    font_subtitle = Font(name='Arial', size=10, bold=True, color='334155')
    font_meta = Font(name='Arial', size=9, bold=True, color='475569')
    font_header = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    font_data = Font(name='Arial', size=9)
    
    fill_header_navy = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    fill_header_blue = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    fill_header_indigo = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
    fill_header_teal = PatternFill(start_color='0D9488', end_color='0D9488', fill_type='solid')
    fill_header_slate = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_success = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
    
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    
    thin_border_side = Side(border_style='thin', color='CBD5E1')
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Estudiantes inscriptos en la carrera
    alumnos_carrera = list(
        Alumno.objects.filter(
            cursadas__comision__plan_estudio__carrera=carrera_obj
        ).select_related('persona').distinct().order_by('persona__apellido', 'persona__nombre')
    )

    # Estructura del plan de estudios de la carrera
    planes_carrera = list(
        PlanEstudio.objects.filter(carrera=carrera_obj).select_related('materia').order_by('anio_carrera', 'id_plan')
    )
    planes_carrera_ids = set(p.pk for p in planes_carrera)
    total_materias = len(planes_carrera)
    total_horas_anuales = sum(p.carga_horaria_anual or 0 for p in planes_carrera)

    anios_en_plan = sorted(list(set(p.anio_carrera for p in planes_carrera if p.anio_carrera and p.anio_carrera > 0)))
    duracion_anios = max(anios_en_plan) if anios_en_plan else 3
    anios_lista = list(range(1, max(duracion_anios, 1) + 1))

    materias_por_anio = defaultdict(list)
    for p in planes_carrera:
        an = p.anio_carrera if p.anio_carrera and p.anio_carrera > 0 else 1
        materias_por_anio[an].append(p)

    # Cargar cursadas y evaluaciones asociadas
    cursadas_qs = Cursada.objects.filter(
        comision__plan_estudio__carrera=carrera_obj,
        alumno__in=alumnos_carrera
    ).select_related(
        'comision__plan_estudio'
    ).prefetch_related(
        'evaluaciones'
    )

    cursadas_by_alumno_plan = {}
    cursadas_by_alumno = defaultdict(list)
    for c in cursadas_qs:
        p_id = c.comision.plan_estudio_id if c.comision else None
        if p_id:
            cursadas_by_alumno_plan[(c.alumno_id, p_id)] = c
        cursadas_by_alumno[c.alumno_id].append(c)

    # Hoja: Libro Matriz (Resumen General)
    ws_matriz = wb.create_sheet(title='LIBRO MATRIZ')
    ws_matriz.views.sheetView[0].showGridLines = True
    
    # Encabezado Institucional y Metadatos de Carrera
    ws_matriz.merge_cells('A1:G1')
    ws_matriz['A1'] = "INSTITUTO SUPERIOR DE FORMACIÓN TÉCNICA N° 188 — LIBRO MATRIZ"
    ws_matriz['A1'].font = font_title
    ws_matriz['A1'].alignment = align_center

    res_str = carrera_obj.resolucion_vigente or 'Sin resolución'
    if carrera_obj.resolucion_anterior:
        res_str += f" (Ant: {carrera_obj.resolucion_anterior})"

    ws_matriz.merge_cells('A2:G2')
    ws_matriz['A2'] = f"CARRERA: {carrera_obj.nombre_carrera} (Código: {carrera_obj.codigo_carrera}) — RESOLUCIÓN: {res_str}"
    ws_matriz['A2'].font = font_subtitle
    ws_matriz['A2'].alignment = align_center

    ws_matriz.merge_cells('A3:G3')
    ws_matriz['A3'] = f"CARACTERÍSTICAS DEL PLAN: {duracion_anios} Años de Duración | {total_materias} Asignaturas | Carga Horaria Total: {total_horas_anuales} hs anuales"
    ws_matriz['A3'].font = font_meta
    ws_matriz['A3'].alignment = align_center

    headers_matriz = ['Alumno', 'DNI', 'Libro', 'Folio / Legajo', 'Cohorte', 'Egresado', 'Plan de Estudio']
    for col_num, h_text in enumerate(headers_matriz, 1):
        cell = ws_matriz.cell(row=5, column=col_num, value=h_text)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = align_center
        cell.border = border_cell

    row_curr = 6
    for idx, al in enumerate(alumnos_carrera, 1):
        p = al.persona
        curs_al = cursadas_by_alumno.get(al.pk, [])

        anios_lectivos = [c.comision.anio_lectivo for c in curs_al if c.comision and c.comision.anio_lectivo]
        cohorte = str(min(anios_lectivos)) if anios_lectivos else (str(al.fecha_ingreso.year) if getattr(al, 'fecha_ingreso', None) else "—")

        # Determinar egreso: todas las materias del plan aprobadas
        aprobadas_ids = set(
            c.comision.plan_estudio_id for c in curs_al 
            if c.comision and c.comision.plan_estudio_id in planes_carrera_ids 
            and c.situacion_final in ('Promocionado', 'Final')
        )
        es_egresado = "SI" if (total_materias > 0 and planes_carrera_ids.issubset(aprobadas_ids)) else "NO"
        
        libro_num = (idx // 100) + 1
        folio_num = (idx % 100) or 100

        vals = [
            f"{p.apellido}, {p.nombre}",
            p.dni,
            libro_num,
            al.legajo or f"F-{folio_num}",
            cohorte,
            es_egresado,
            carrera_obj.resolucion_vigente or carrera_obj.codigo_carrera
        ]
        
        bg_fill = fill_zebra if idx % 2 == 0 else PatternFill(fill_type=None)
        if es_egresado == "SI":
            bg_fill = fill_success

        for col_num, val in enumerate(vals, 1):
            cell = ws_matriz.cell(row=row_curr, column=col_num, value=val)
            cell.font = font_data
            if bg_fill.fill_type: cell.fill = bg_fill
            cell.border = border_cell
            cell.alignment = align_center if col_num in [2, 3, 4, 5, 6] else align_left
        row_curr += 1

    # Ajustar anchos de columnas
    for col in ws_matriz.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_matriz.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Hojas de calificaciones por año de cursada
    colores_headers = [fill_header_blue, fill_header_indigo, fill_header_navy, fill_header_teal, fill_header_slate]

    for i_anio, anio in enumerate(anios_lista):
        sheet_title = f"{anio}° Año"
        ws_anio = wb.create_sheet(title=sheet_title)
        ws_anio.views.sheetView[0].showGridLines = True
        
        planes_anio = materias_por_anio.get(anio, [])
        total_cols = max(3 + len(planes_anio), 4)

        ws_anio.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        ws_anio.cell(row=1, column=1, value=f"PLANILLA DE CALIFICACIONES — {anio}° AÑO ({carrera_obj.nombre_carrera})").font = font_title
        ws_anio.cell(row=1, column=1).alignment = align_center

        fill_header_anio = colores_headers[i_anio % len(colores_headers)]

        # Headers
        headers_anio = ['Estudiante', 'DNI', 'Folio']
        for p in planes_anio:
            headers_anio.append(p.materia.nombre_materia)

        for col_num, h_text in enumerate(headers_anio, 1):
            cell = ws_anio.cell(row=3, column=col_num, value=h_text)
            cell.font = font_header
            cell.fill = fill_header_anio
            cell.alignment = align_center
            cell.border = border_cell

        row_curr = 4
        for idx, al in enumerate(alumnos_carrera, 1):
            p = al.persona
            folio_num = al.legajo or f"{idx}"
            row_data = [f"{p.apellido}, {p.nombre}", p.dni, folio_num]

            # Buscar notas de cada materia en este año desde la memoria
            for plan_it in planes_anio:
                curs = cursadas_by_alumno_plan.get((al.pk, plan_it.pk))
                if curs:
                    evals_validas = [ev for ev in curs.evaluaciones.all() if ev.nota is not None]
                    if evals_validas:
                        evals_validas.sort(key=lambda ev: ev.fecha or datetime.date.min, reverse=True)
                        row_data.append(float(evals_validas[0].nota))
                    elif curs.situacion_final:
                        row_data.append(curs.situacion_final)
                    else:
                        row_data.append("")
                else:
                    row_data.append("")

            bg_fill = fill_zebra if idx % 2 == 0 else PatternFill(fill_type=None)
            for col_num, val in enumerate(row_data, 1):
                cell = ws_anio.cell(row=row_curr, column=col_num, value=val)
                cell.font = font_data
                if bg_fill.fill_type: cell.fill = bg_fill
                cell.border = border_cell
                cell.alignment = align_center if col_num > 1 else align_left
            row_curr += 1

        for col in ws_anio.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_anio.column_dimensions[col_letter].width = max(max_len + 3, 11)

    # Hoja: Egresados
    ws_egr = wb.create_sheet(title='Egresados')
    ws_egr.views.sheetView[0].showGridLines = True
    
    ws_egr.merge_cells('A1:F1')
    ws_egr.cell(row=1, column=1, value=f"PLANILLA DE EGRESADOS — {carrera_obj.nombre_carrera} ({duracion_anios} Años)").font = font_title
    ws_egr.cell(row=1, column=1).alignment = align_center

    headers_egr = ['Estudiante', 'DNI', 'Libro', 'Folio', 'Egreso', 'Año de Egreso']
    for col_num, h_text in enumerate(headers_egr, 1):
        cell = ws_egr.cell(row=3, column=col_num, value=h_text)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = align_center
        cell.border = border_cell

    row_curr = 4
    for idx, al in enumerate(alumnos_carrera, 1):
        p = al.persona
        curs_al = cursadas_by_alumno.get(al.pk, [])

        aprobadas = [
            c for c in curs_al 
            if c.comision and c.comision.plan_estudio_id in planes_carrera_ids 
            and c.situacion_final in ('Promocionado', 'Final')
        ]
        aprobadas_ids = set(c.comision.plan_estudio_id for c in aprobadas)
        es_egr = "SI" if (total_materias > 0 and planes_carrera_ids.issubset(aprobadas_ids)) else "NO"
        
        if es_egr == "SI":
            anios_aprobadas = [c.comision.anio_lectivo for c in aprobadas if c.comision.anio_lectivo]
            anio_egreso = str(max(anios_aprobadas)) if anios_aprobadas else str(datetime.date.today().year)
        else:
            anio_egreso = ""

        libro_num = (idx // 100) + 1
        folio_num = al.legajo or f"{idx}"

        vals = [f"{p.apellido}, {p.nombre}", p.dni, libro_num, folio_num, es_egr, anio_egreso]
        bg_fill = fill_success if es_egr == "SI" else (fill_zebra if idx % 2 == 0 else PatternFill(fill_type=None))

        for col_num, val in enumerate(vals, 1):
            cell = ws_egr.cell(row=row_curr, column=col_num, value=val)
            cell.font = font_data
            if bg_fill.fill_type: cell.fill = bg_fill
            cell.border = border_cell
            cell.alignment = align_center if col_num > 1 else align_left
        row_curr += 1

    for col in ws_egr.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_egr.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
